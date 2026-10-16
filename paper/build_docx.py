"""Generates PAPER_IEEE.docx from PAPER_DRAFT.md (the Markdown file is the source of truth).
Usage (from the repository root): python build_docx.py
Requires: python-docx
"""
import re
from docx import Document
from docx.shared import Pt, Mm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

SRC = "PAPER_DRAFT.md"
OUT = "PAPER_IEEE.docx"
FONT = "Times New Roman"
COLW = (210 - 29 - 4.2) / 2  # column width in mm


def set_font(run, size=None, bold=None, italic=None, name=FONT, sup=None, sub=None, small_caps=None):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rf = rpr.find(qn('w:rFonts'))
    if rf is None:
        rf = OxmlElement('w:rFonts')
        rpr.append(rf)
    for a in ('w:ascii', 'w:hAnsi', 'w:cs', 'w:eastAsia'):
        rf.set(qn(a), name)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if sup:
        run.font.superscript = True
    if sub:
        run.font.subscript = True
    if small_caps:
        run.font.small_caps = True


def math_runs(tex):
    """Convert the small LaTeX subset used in the draft into (text, kind) segments."""
    t = tex
    t = t.replace(r"\dot{V}", "V\u0307").replace(r"\mathrm{VDOT}", "VDOT").replace(r"\cdot", "\u00b7")
    t = t.replace(r"\,", "\u2009").replace(r"\;", " ")
    t = re.sub(r"\\tag\{[^}]*\}", "", t)
    segs = []
    i = 0
    while i < len(t):
        c = t[i]
        if c in "^_" and i + 1 < len(t):
            kind = 'sup' if c == '^' else 'sub'
            if t[i + 1] == '{':
                j = t.index('}', i + 2)
                segs.append((t[i + 2:j], kind))
                i = j + 1
            else:
                segs.append((t[i + 1], kind))
                i += 2
        else:
            if segs and segs[-1][1] == 'n':
                segs[-1] = (segs[-1][0] + c, 'n')
            else:
                segs.append((c, 'n'))
            i += 1
    return segs


INLINE = re.compile(r"(\*\*.+?\*\*|\*[^*\s][^*]*?\*|`[^`]+`|\$[^$]+\$)")


def add_inline(par, text, size, base_bold=False, base_italic=False):
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            set_font(par.add_run(part[2:-2]), size, bold=True, italic=base_italic)
        elif part.startswith("`") and part.endswith("`"):
            set_font(par.add_run(part[1:-1]), size - 1, name="Courier New", bold=base_bold)
        elif part.startswith("$") and part.endswith("$"):
            for seg, k in math_runs(part[1:-1]):
                set_font(par.add_run(seg), size, bold=base_bold, italic=True if k == 'n' else None,
                         sup=(k == 'sup'), sub=(k == 'sub'))
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            set_font(par.add_run(part[1:-1]), size, bold=base_bold, italic=True)
        else:
            set_font(par.add_run(part), size, bold=base_bold, italic=base_italic)


def fmt(par, align=WD_ALIGN_PARAGRAPH.JUSTIFY, before=0, after=0, left=None, first=None, keep_next=False):
    pf = par.paragraph_format
    par.alignment = align
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    if left is not None:
        pf.left_indent = Mm(left)
    if first is not None:
        pf.first_line_indent = Mm(first)
    pf.keep_with_next = keep_next


def cols(section, n, space_mm=4.2):
    sp = section._sectPr
    c = sp.find(qn('w:cols'))
    if c is None:
        c = OxmlElement('w:cols')
        sp.append(c)
    c.set(qn('w:num'), str(n))
    c.set(qn('w:space'), str(int(space_mm * 56.7)))


def page_setup(section):
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.top_margin = Mm(19)
    section.bottom_margin = Mm(25)
    section.left_margin = Mm(14.5)
    section.right_margin = Mm(14.5)


def add_page_number(section):
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for typ, txt in (('begin', None), (None, 'PAGE'), ('end', None)):
        r = p.add_run()
        set_font(r, 9)
        if typ:
            f = OxmlElement('w:fldChar')
            f.set(qn('w:fldCharType'), typ)
            r._element.append(f)
        else:
            it = OxmlElement('w:instrText')
            it.set(qn('xml:space'), 'preserve')
            it.text = txt
            r._element.append(it)


def shade(cell, color):
    tcPr = cell._element.get_or_add_tcPr()
    sh = OxmlElement('w:shd')
    sh.set(qn('w:val'), 'clear')
    sh.set(qn('w:color'), 'auto')
    sh.set(qn('w:fill'), color)
    tcPr.append(sh)


def borders(table):
    tblPr = table._element.tblPr
    b = OxmlElement('w:tblBorders')
    for e in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        x = OxmlElement(f'w:{e}')
        x.set(qn('w:val'), 'single')
        x.set(qn('w:sz'), '4')
        x.set(qn('w:space'), '0')
        x.set(qn('w:color'), '000000')
        b.append(x)
    tblPr.append(b)


def cell_margins(table, tw=40):
    tblPr = table._element.tblPr
    m = OxmlElement('w:tblCellMar')
    for e in ('top', 'left', 'bottom', 'right'):
        x = OxmlElement(f'w:{e}')
        x.set(qn('w:w'), str(tw))
        x.set(qn('w:type'), 'dxa')
        m.append(x)
    tblPr.append(m)


def add_table(doc, rows):
    n = len(rows[0])
    t = doc.add_table(rows=len(rows), cols=n)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    borders(t)
    cell_margins(t)
    t.autofit = False
    lens = [max(len(re.sub(r"[`*$]", "", r[i])) for r in rows) for i in range(n)]
    lens = [min(max(l, 10), 60) for l in lens]
    tot = sum(lens)
    for ri, r in enumerate(rows):
        for ci, txt in enumerate(r):
            c = t.cell(ri, ci)
            c.width = Mm(COLW * lens[ci] / tot)
            p = c.paragraphs[0]
            fmt(p, WD_ALIGN_PARAGRAPH.LEFT)
            add_inline(p, txt, 7.5, base_bold=(ri == 0))
            if ri == 0:
                shade(c, 'D9D9D9')
    trPr = t.rows[0]._tr.get_or_add_trPr()
    h = OxmlElement('w:tblHeader')
    h.set(qn('w:val'), 'true')
    trPr.append(h)
    sp = doc.add_paragraph()
    fmt(sp, after=4)
    set_font(sp.add_run(""), 4)
    return t


txt = open(SRC, encoding='utf8').read()
txt = re.sub(r"<!--.*?-->\n*", "", txt, flags=re.S)
lines = txt.split("\n")

doc = Document()
st = doc.styles['Normal']
st.font.name = FONT
st.font.size = Pt(10)
st.element.rPr.rFonts.set(qn('w:eastAsia'), FONT)
sec = doc.sections[0]
page_setup(sec)
add_page_number(sec)


def nxt(k):
    while k < len(lines) and not lines[k].strip():
        k += 1
    return k


i = nxt(0)
assert lines[i].startswith("# ")
title = lines[i][2:].strip()
i += 1
p = doc.add_paragraph()
fmt(p, WD_ALIGN_PARAGRAPH.CENTER, after=4)
set_font(p.add_run(title), 24)
i = nxt(i)
p = doc.add_paragraph()
fmt(p, WD_ALIGN_PARAGRAPH.CENTER, after=8)
add_inline(p, lines[i].strip(), 10)
i += 1
while not lines[i].startswith("---"):
    if lines[i].strip():
        p = doc.add_paragraph()
        fmt(p, WD_ALIGN_PARAGRAPH.CENTER, after=2)
        add_inline(p, lines[i].strip(), 10)
    i += 1
i += 1
new = doc.add_section(WD_SECTION.CONTINUOUS)
page_setup(new)
cols(new, 2)

refs_mode = False
DASH = "\u2014"
while i < len(lines):
    ln = lines[i].rstrip()
    if not ln.strip() or ln.strip() == "---":
        i += 1
        continue
    if ln.startswith("**Abstract**"):
        body = ln.split(DASH, 1)[1]
        p = doc.add_paragraph()
        fmt(p, after=4)
        set_font(p.add_run("Abstract" + DASH), 9, bold=True, italic=True)
        add_inline(p, body, 9, base_bold=True)
        i += 1
        continue
    if ln.startswith("**Index Terms**"):
        body = ln.split(DASH, 1)[1]
        p = doc.add_paragraph()
        fmt(p, after=8)
        set_font(p.add_run("Index Terms" + DASH), 9, bold=True, italic=True)
        add_inline(p, body, 9, base_bold=True)
        i += 1
        continue
    if ln.startswith("## "):
        h = ln[3:].strip()
        refs_mode = h.lower() == "references"
        p = doc.add_paragraph(style='Heading 1')
        fmt(p, WD_ALIGN_PARAGRAPH.CENTER, before=10, after=4, keep_next=True)
        r = p.add_run(h)
        set_font(r, 10, bold=False, italic=False, small_caps=True)
        r.font.color.rgb = RGBColor(0, 0, 0)
        i += 1
        continue
    if ln.startswith("### "):
        p = doc.add_paragraph(style='Heading 2')
        fmt(p, WD_ALIGN_PARAGRAPH.LEFT, before=6, after=3, keep_next=True)
        r = p.add_run(ln[4:].strip())
        set_font(r, 10, italic=True, bold=False)
        r.font.color.rgb = RGBColor(0, 0, 0)
        i += 1
        continue
    if ln.strip() == "[FIGURE TO BE GENERATED]":
        t = doc.add_table(rows=1, cols=1)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        borders(t)
        cell_margins(t, 120)
        c = t.cell(0, 0)
        c.width = Mm(COLW)
        shade(c, 'F2F2F2')
        p = c.paragraphs[0]
        fmt(p, WD_ALIGN_PARAGRAPH.CENTER, before=18, after=18)
        set_font(p.add_run("[FIGURE TO BE GENERATED]"), 9, bold=True)
        i += 1
        continue
    if ln.startswith("**Fig."):
        p = doc.add_paragraph()
        fmt(p, WD_ALIGN_PARAGRAPH.CENTER, before=3, after=8)
        m = re.match(r"\*\*(Fig\. \d+\.)\*\*\s*(.*)", ln)
        set_font(p.add_run(m.group(1) + " "), 8)
        add_inline(p, m.group(2), 8)
        i += 1
        continue
    if ln.startswith("**TABLE"):
        m = re.match(r"\*\*TABLE ([IVX]+)\.\s*(.*?)\*\*$", ln)
        p = doc.add_paragraph()
        fmt(p, WD_ALIGN_PARAGRAPH.CENTER, before=4, after=1, keep_next=True)
        set_font(p.add_run("TABLE " + m.group(1)), 8)
        p2 = doc.add_paragraph()
        fmt(p2, WD_ALIGN_PARAGRAPH.CENTER, after=3, keep_next=True)
        add_inline(p2, m.group(2), 8)
        i += 1
        continue
    if ln.startswith("|"):
        rows = []
        while i < len(lines) and lines[i].startswith("|"):
            cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
            if not all(re.fullmatch(r":?-{3,}:?", c) for c in cells):
                rows.append(cells)
            i += 1
        add_table(doc, rows)
        continue
    if ln.startswith("$$"):
        eq = ln.strip().strip("$").strip()
        num = re.search(r"\\tag\{(\d+)\}", eq).group(1)
        p = doc.add_paragraph()
        fmt(p, WD_ALIGN_PARAGRAPH.LEFT, before=3, after=3)
        ts = p.paragraph_format.tab_stops
        ts.add_tab_stop(Mm(COLW / 2), WD_TAB_ALIGNMENT.CENTER)
        ts.add_tab_stop(Mm(COLW), WD_TAB_ALIGNMENT.RIGHT)
        set_font(p.add_run("\t"), 9)
        for seg, k in math_runs(eq):
            set_font(p.add_run(seg), 9, italic=True if k == 'n' else None, sup=(k == 'sup'), sub=(k == 'sub'))
        set_font(p.add_run("\t(" + num + ")"), 10)
        i += 1
        continue
    m = re.match(r"^\[(\d+)\] (.*)", ln)
    if refs_mode and m:
        p = doc.add_paragraph()
        fmt(p, WD_ALIGN_PARAGRAPH.JUSTIFY, after=2, left=7, first=-7)
        p.paragraph_format.tab_stops.add_tab_stop(Mm(7))
        set_font(p.add_run("[" + m.group(1) + "]\t"), 8)
        add_inline(p, m.group(2), 8)
        i += 1
        continue
    if ln.startswith("- "):
        p = doc.add_paragraph()
        fmt(p, WD_ALIGN_PARAGRAPH.JUSTIFY, after=2, left=4)
        add_inline(p, ln[2:], 10)
        i += 1
        continue
    p = doc.add_paragraph()
    fmt(p, WD_ALIGN_PARAGRAPH.JUSTIFY, after=4, first=4)
    add_inline(p, ln.strip(), 10)
    i += 1

doc.core_properties.title = title
doc.core_properties.author = "Priscillia Lovemel Candra"
doc.save(OUT)
print("saved", OUT)
